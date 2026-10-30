import io
import zipfile
from contextlib import closing

import pytest
from docx import Document
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from pi.api.routes.artifacts import answers, router
from pi.content import artifacts as a
from pi.conversation.store import Store
from pi.core.errors import install


@pytest.fixture
def api(tmp_path):
    with closing(Store(tmp_path / "pi.db")) as store:
        state = {"owner": False, "unknown": False}

        def authorize():
            if not state["owner"]:
                raise HTTPException(403)

        def resolve(db, session_id):
            return None if state["unknown"] else {"memoryDisabled": False, "harnessDisabled": False}

        app = FastAPI()
        install(app)
        app.include_router(router(lambda: store, authorize, resolve))
        app.include_router(answers(lambda: store, authorize, resolve))
        with TestClient(app) as client:
            yield store, client, state, resolve


def test_xlsx_literal_cells_download_auth_and_version(api):
    store, client, state, _ = api
    table = {
        "kind": "table",
        "columns": ["Name", "Value"],
        "rows": [
            ["Unicode \u05e9\u05dc\u05d5\u05dd", '=WEBSERVICE("https://invalid")'],
            ["Count", "0012"],
            ["Markup", "<script>&"],
        ],
    }
    record = a.create(store, a.Create(title="../../CON", content=table))
    path = f"/artifacts/{record['id']}/download?format=xlsx&version=1"
    assert client.get(path).status_code == 403
    state["owner"] = True
    response = client.get(path)
    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["x-content-type-options"] == "nosniff"
    assert response.headers["content-disposition"] == 'attachment; filename="artifact-CON-v1.xlsx"'
    book = load_workbook(io.BytesIO(response.content))
    assert list(book.active.values) == [tuple(table["columns"]), *map(tuple, table["rows"])]
    assert book.active["B2"].data_type == "s" and book.active.freeze_panes == "A2"
    book.close()
    with zipfile.ZipFile(io.BytesIO(response.content)) as package:
        assert not any("externalLinks" in n or "vbaProject" in n for n in package.namelist())
        assert b"<f>" not in package.read("xl/worksheets/sheet1.xml")
    assert client.get(path.replace("version=1", "version=99")).status_code == 422
    assert client.get(path.replace("format=xlsx", "format=exe")).status_code == 422
    document = Document(io.BytesIO(client.get(path.replace("format=xlsx", "format=docx")).content))
    assert [[c.text for c in row.cells] for row in document.tables[0].rows] == [
        table["columns"],
        *table["rows"],
    ]


def test_docx_editable_headings_and_source_privacy(api):
    store, client, state, resolve = api
    session = store.create_session()
    turn = store.start_turn(session)
    message = store.complete_turn(turn, "# Report\nHello & <world>\n\u05e9\u05dc\u05d5\u05dd")
    record = a.create(
        store,
        a.FromMessage(title="Report", sessionId=session, messageId=message["id"]),
        resolve=resolve,
    )
    state["owner"] = True
    path = f"/artifacts/{record['id']}/download?format=docx"
    response = client.get(path)
    assert response.status_code == 200
    document = Document(io.BytesIO(response.content))
    # One markdown paragraph: a single line break reads as a space, as in the chat.
    assert [p.text for p in document.paragraphs] == [
        "Report",
        "Hello & <world> \u05e9\u05dc\u05d5\u05dd",
    ]
    assert document.paragraphs[0].style.name == "Heading 1"
    state["unknown"] = True
    assert client.get(path).status_code != 200
    assert b"Hello" not in client.get(path).content


def test_native_html_stays_attachment_and_office_controls_are_rejected(api):
    store, client, state, _ = api
    state["owner"] = True
    record = a.create(
        store, a.Create(title="Page", content={"kind": "html", "text": "<script>x</script>"})
    )
    response = client.get(f"/artifacts/{record['id']}/download")
    assert response.content == b"<script>x</script>"
    assert response.headers["content-type"].startswith("text/plain")
    assert response.headers["content-disposition"].endswith('.html.txt"')
    record = a.create(
        store, a.Create(title="Control", content={"kind": "markdown", "text": "bad\x00text"})
    )
    response = client.get(f"/artifacts/{record['id']}/download?format=docx")
    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "export_control"


def test_pdf_and_pptx_downloads_keep_citations_and_a_readable_name(api):
    store, client, state, resolve = api
    state["owner"] = True
    session = store.create_session()
    message = store.complete_turn(
        store.start_turn(session),
        "## Findings\n\n- One [source](https://example.com/a)\n- Two",
        citations=[{"id": "c1", "label": "Example", "href": "https://example.com/c"}],
    )
    title = "Отчёт"
    record = a.create(
        store,
        a.FromMessage(title=title, sessionId=session, messageId=message["id"]),
        resolve=resolve,
    )
    response = client.get(f"/artifacts/{record['id']}/download?format=pdf")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"] == (
        "attachment; filename=\"artifact-v1.pdf\"; filename*=UTF-8''%D0%9E%D1%82%D1%87%D1%91%D1%82-v1.pdf"
    )
    text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(response.content)).pages)
    assert "Findings" in text and "Supplied source references" in text
    response = client.get(f"/artifacts/{record['id']}/download?format=pptx")
    assert response.headers["content-type"].endswith("presentationml.presentation")
    deck = Presentation(io.BytesIO(response.content))
    assert deck.slides[0].shapes[0].text_frame.text == "Findings"


def test_answer_download_checks_source_privacy_and_owner(api):
    store, client, state, _ = api
    session = store.create_session()
    message = store.complete_turn(
        store.start_turn(session), "# Plan\n\n| A | B |\n|--|--|\n| 1 | 2 |"
    )
    path = f"/sessions/{session}/messages/{message['id']}/download"
    assert client.get(path + "?format=docx").status_code == 403
    state["owner"] = True
    response = client.get(path + "?format=xlsx")
    assert response.status_code == 200
    assert response.headers["content-disposition"] == 'attachment; filename="Plan.xlsx"'
    assert response.headers["cache-control"] == "no-store"
    assert list(load_workbook(io.BytesIO(response.content)).active.values) == [
        ("A", "B"),
        ("1", "2"),
    ]
    document = Document(io.BytesIO(client.get(path + "?format=docx").content))
    assert document.paragraphs[0].text == "Plan"
    assert client.get(path).status_code == 422  # A format is required.
    assert client.get(path + "?format=native").status_code == 422
    other = store.create_session()
    missing = client.get(f"/sessions/{other}/messages/{message['id']}/download?format=pdf")
    assert missing.status_code == 404 and missing.json()["detail"]["code"] == "not_found"
    assert client.get(f"/sessions/{session}/messages/msg_x/download?format=pdf").status_code == 422
    asked = store.append_message(session, "user", "Please export me")
    refused = client.get(f"/sessions/{session}/messages/{asked['id']}/download?format=pdf")
    assert refused.status_code == 422 and refused.json()["detail"]["code"] == "not_an_answer"
    state["unknown"] = True
    hidden = client.get(path + "?format=pdf")
    assert hidden.status_code == 409 and hidden.json()["detail"]["code"] == "privacy_unknown"
    assert b"Plan" not in hidden.content
