"""DOCX, PDF, PPTX and XLSX exports reopened with each format's own reader."""

import io
import zipfile

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from pi.content import export_documents, export_pdf, export_slides, office_exports
from pi.content import markdown_blocks as m

SOURCE = (
    "# Quarterly report\n\n"
    "Intro with **bold**, *italic*, `code`, a [site](https://example.com/site) "
    "and [trap](javascript:alert(1)).\n\n"
    "## Plan\n\n"
    "1. First step\n   - nested idea\n2. Second step\n\n"
    "> Quoted note\n\n"
    "```python\ndef double(x):\n    return x * 2\n```\n\n"
    "### Numbers\n\n"
    "| Name | Value |\n|:--|--:|\n| Alpha | =1+2 |\n| Beta | [docs](https://example.com/docs) |\n"
)
CITATIONS = [{"id": "c1", "label": "Source one", "href": "https://example.com/source"}]


@pytest.fixture(scope="module")
def blocks():
    return m.from_markdown(SOURCE) + m.citation_blocks(CITATIONS)


def external_links(package):
    with zipfile.ZipFile(io.BytesIO(package)) as files:
        rels = [n for n in files.namelist() if n.endswith(".rels")]
        text = b"".join(files.read(n) for n in rels) + b"".join(
            files.read(n) for n in files.namelist() if n.endswith(".xml")
        )
    assert b"javascript" not in text.lower()
    return text


def test_docx_keeps_headings_lists_quotes_code_tables_and_safe_links(blocks):
    package = export_documents.docx(blocks, "Quarterly report")
    document = Document(io.BytesIO(package))
    paragraphs = [(p.style.name, p.text) for p in document.paragraphs if p.text]
    assert paragraphs[:4] == [
        ("Heading 1", "Quarterly report"),
        ("Normal", "Intro with bold, italic, code, a site and trap."),
        ("Heading 2", "Plan"),
        ("Normal", "1.\tFirst step"),
    ]
    assert ("Normal", "◦\tnested idea") in paragraphs
    assert ("Quote", "Quoted note") in paragraphs
    assert ("Normal", "def double(x):\n    return x * 2") in paragraphs
    assert ("Heading 3", "Numbers") in paragraphs
    assert ("Normal", "1.\tSource one https://example.com/source") in paragraphs
    (table,) = document.tables
    assert [[c.text for c in row.cells] for row in table.rows] == [
        ["Name", "Value"],
        ["Alpha", "=1+2"],
        ["Beta", "docs"],
    ]
    targets = {r.target_ref for r in document.part.rels.values() if r.is_external}
    assert targets == {
        "https://example.com/site",
        "https://example.com/docs",
        "https://example.com/source",
    }
    assert document.core_properties.title == "Quarterly report"
    external_links(package)


def test_pdf_holds_the_text_and_only_safe_link_annotations(blocks):
    reader = PdfReader(io.BytesIO(export_pdf.pdf(blocks, "Quarterly report")))
    text = "\n".join(page.extract_text() for page in reader.pages)
    for expected in (
        "Quarterly report",
        "Intro with bold, italic, code, a site and trap.",
        "First step",
        "Quoted note",
        "return x * 2",
        "Alpha",
        "=1+2",
        "Supplied source references",
    ):
        assert expected in text
    uris = [
        annotation.get_object()["/A"]["/URI"]
        for page in reader.pages
        for annotation in page.get("/Annots") or []
    ]
    assert sorted(uris) == [
        "https://example.com/docs",
        "https://example.com/site",
        "https://example.com/source",
    ]
    assert reader.metadata.title == "Quarterly report"


@pytest.mark.skipif(not (export_pdf.DEJAVU / "DejaVuSans.ttf").is_file(), reason="no DejaVu")
def test_pdf_draws_cyrillic_and_hebrew_with_dejavu():
    blocks = m.from_markdown("Привет שלום")
    text = PdfReader(io.BytesIO(export_pdf.pdf(blocks, "t"))).pages[0].extract_text()
    assert "Привет" in text


def test_pdf_page_and_layout_limits(monkeypatch):
    monkeypatch.setattr(export_pdf, "MAX_PAGES", 2)
    with pytest.raises(m.ExportLimit) as caught:
        export_pdf.pdf(m.from_markdown("\n\n".join(["Paragraph"] * 400)), "Long")
    assert caught.value.code == "pages"
    tall = m.from_table(["Only"], [["word " * 6000]])
    with pytest.raises(m.ExportLimit) as caught:
        export_pdf.pdf(tall, "Tall")
    assert caught.value.code == "layout"


def test_pdf_markup_in_text_stays_text():
    blocks = m.from_markdown('<img src="https://example.com/x.png"/> <a href="javascript:x">y</a>')
    package = export_pdf.pdf(blocks, "Markup")
    text = PdfReader(io.BytesIO(package)).pages[0].extract_text()
    assert '<img src="https://example.com/x.png"/>' in text
    assert b"/URI" not in package and b"/XObject" not in package


def test_pptx_slides_follow_headings_and_tables_get_their_own(blocks):
    deck = Presentation(io.BytesIO(export_slides.pptx(blocks, "Quarterly report")))

    def contents(slide):
        return [
            shape.text_frame.text
            if shape.has_text_frame
            else [[c.text for c in row.cells] for row in shape.table.rows]
            for shape in slide.shapes
        ]

    slides = [contents(slide) for slide in deck.slides]
    assert slides[0] == ["Quarterly report", "Intro with bold, italic, code, a site and trap."]
    assert slides[1][0] == "Plan" and "1.\tFirst step" in slides[1][1]
    assert "def double(x):\x0b    return x * 2" in slides[1][1]
    assert slides[2] == [
        "Plan (continued)",
        [["Name", "Value"], ["Alpha", "=1+2"], ["Beta", "docs"]],
    ]
    assert slides[3][0] == "Supplied source references"


def test_pptx_long_tables_repeat_the_header_and_slides_are_capped(monkeypatch):
    rows = [[f"r{i}", str(i)] for i in range(25)]
    deck = Presentation(io.BytesIO(export_slides.pptx(m.from_table(["A", "B"], rows), "Big")))
    tables = [s.shapes[1].table for s in deck.slides]
    assert [len(t.rows) for t in tables] == [11, 11, 6]
    assert all(t.cell(0, 0).text == "A" for t in tables)
    assert [s.shapes[0].text_frame.text for s in deck.slides] == [
        "Big",
        "Big (continued)",
        "Big (continued)",
    ]
    monkeypatch.setattr(export_slides, "MAX_SLIDES", 2)
    with pytest.raises(m.ExportLimit) as caught:
        export_slides.pptx(m.from_table(["A", "B"], rows), "Big")
    assert caught.value.code == "slides"


def test_xlsx_puts_each_markdown_table_on_a_sheet_named_by_its_heading():
    blocks = m.from_markdown(
        "## Costs: 2026/Q1\n\n| Item | Cost |\n|--|--|\n| Rent | =SUM(A1) |\n\n"
        "| Link |\n|--|\n| [site](https://example.com) |\n"
    )
    book = load_workbook(io.BytesIO(export_documents.xlsx(blocks, "Costs")))
    assert book.sheetnames == ["Costs 2026 Q1", "Table 2"]
    first, second = book.worksheets
    assert list(first.values) == [("Item", "Cost"), ("Rent", "=SUM(A1)")]
    assert first["B2"].data_type == "s"
    assert second["A2"].hyperlink.target == "https://example.com"
    with pytest.raises(m.ExportLimit) as caught:
        export_documents.xlsx(m.from_markdown("No table here."), "None")
    assert caught.value.code == "no_table"


def test_names_are_safe_and_keep_a_readable_unicode_name():
    assert office_exports.names("../../CON") == ("artifact-CON", "artifact-CON")
    assert office_exports.names("Отчёт 2026") == (
        "2026",
        "Отчёт-2026",
    )
    assert office_exports.names("‮cod.exe\\..\\x") == ("cod-exe-x", "cod-exe-x")
    assert office_exports.names("שלום")[0] == "artifact"
    assert office_exports.disposition("2026.pdf", "От-2026.pdf") == (
        "attachment; filename=\"2026.pdf\"; filename*=UTF-8''%D0%9E%D1%82-2026.pdf"
    )
    assert office_exports.disposition("a.pdf", "a.pdf") == 'attachment; filename="a.pdf"'


def test_worker_makes_every_format_and_refuses_with_fixed_codes():
    data = {"kind": "markdown", "text": SOURCE}
    for format, magic in office_exports.MAGIC.items():
        assert office_exports.render(data, format, title="Report").startswith(magic)
    with pytest.raises(office_exports.ExportError) as caught:
        office_exports.render({"kind": "markdown", "text": "- " * 200 + "x"}, "docx", title="t")
    assert (caught.value.code, caught.value.status) == ("export_too_deep", 422)
    assert "12 levels" in caught.value.message


def test_refusals_before_any_worker_starts(monkeypatch):
    def never(*args, **kwargs):
        raise AssertionError("no worker for refused content")

    monkeypatch.setattr(office_exports.subprocess, "run", never)
    for data, format, code in (
        ({"kind": "markdown", "text": "bad\x00text"}, "pdf", "export_control"),
        ({"kind": "markdown", "text": "\ud800"}, "docx", "export_control"),
        ({"kind": "html", "text": "<p>x</p>"}, "pdf", "unsupported_format"),
        ({"kind": "code", "text": "x", "language": "py"}, "xlsx", "unsupported_format"),
        ({"kind": "markdown", "text": "x"}, "exe", "unsupported_format"),
        ({"kind": "markdown", "text": "x" * (3 * 1024 * 1024)}, "pdf", "export_too_long"),
    ):
        with pytest.raises(office_exports.ExportError) as caught:
            office_exports.render(data, format, title="t")
        assert (caught.value.code, caught.value.status) == (code, 422)


def test_export_limit_matches_the_worker():
    from pi.content import _export_worker

    assert _export_worker.MAX_INPUT == office_exports.MAX_INPUT


def test_a_worker_that_overruns_its_time_is_reported_unavailable(monkeypatch):
    monkeypatch.setattr(office_exports, "TIMEOUT", 0.01)
    with pytest.raises(office_exports.ExportError) as caught:
        office_exports.render({"kind": "markdown", "text": "x"}, "pdf", title="t")
    assert (caught.value.code, caught.value.status) == ("export_unavailable", 503)
