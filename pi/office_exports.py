"""Bounded Office exports from validated inert artifact content; no imports or fetches."""

import io
import re


def _text(value):
    if any(ord(char) < 32 and char not in "\t\n\r" for char in value):
        raise ValueError("Office exports cannot encode control characters.")
    if any(0xD800 <= ord(char) <= 0xDFFF or ord(char) in (0xFFFE, 0xFFFF) for char in value):
        raise ValueError("Office exports require valid XML text.")
    return value


def render(data, text, format):
    output = io.BytesIO()
    if format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        book = Workbook()
        sheet = book.active
        sheet.title = "Data"
        for row_index, row in enumerate([data["columns"], *data["rows"]], 1):
            for column, value in enumerate(row, 1):
                cell = sheet.cell(row_index, column)
                cell.value = _text(value)
                cell.data_type = "s"  # Formula-looking owner/model content stays literal.
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_index == 1:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="243B32")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in range(1, len(data["columns"]) + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 24
        book.save(output)
        book.close()
    elif format == "docx":
        from docx import Document
        from docx.shared import Pt

        lines = _text(text).split("\n")
        if len(lines) > 5000:
            raise ValueError("Document exceeds 5000 paragraphs.")
        document = Document()
        document.core_properties.author = "Conker"
        document.core_properties.last_modified_by = "Conker"
        document.styles["Normal"].font.size = Pt(11)
        fenced = False
        for line in lines:
            if line.startswith("```"):
                fenced = not fenced
                document.add_paragraph(line)
                continue
            heading = re.match(r"^(#{1,6}) (.*)$", line) if not fenced else None
            if heading:
                document.add_heading(heading[2], level=len(heading[1]))
            else:
                document.add_paragraph(line)
        document.save(output)
    else:
        raise ValueError("Unsupported Office format.")
    if output.tell() > 2 * 1024 * 1024:
        raise ValueError("Export exceeds 2 MiB.")
    return output.getvalue()
