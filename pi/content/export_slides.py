"""PowerPoint decks drawn from markdown blocks with python-pptx; nothing is fetched.

Each level 1 or 2 heading starts a slide titled with it, and a horizontal rule starts a new
slide under the same title. Text that would overflow continues on a slide marked
"(continued)", and a table gets slides of its own, ten rows each under a repeated header.
"""

from __future__ import annotations

import io

from pi.content.markdown_blocks import Block, ExportLimit, Span, plain

MAX_SLIDES = 300
LINES = 13  # Body lines a slide holds at 18 points.
CHARACTERS = 95  # Characters a body line holds at 18 points.
CODE_LINES = 24  # Code lines a slide holds at 12 points.
TABLE_ROWS = 10
MONO = "Courier New"
CONTINUED = " (continued)"


def _cell(cell, part, text: str, size: int, href: str | None):
    """One table cell's text; size is in hundredths of a point."""
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE
    from pptx.oxml.ns import qn

    paragraph = cell.find(qn("a:txBody")).find(qn("a:p"))
    for number, line in enumerate(text.split("\n")):
        if number:
            etree.SubElement(paragraph, qn("a:br"))
        if not line:
            continue
        run = etree.SubElement(paragraph, qn("a:r"))
        properties = etree.SubElement(run, qn("a:rPr"), lang="en-US", sz=str(size))
        if href:
            link = etree.SubElement(properties, qn("a:hlinkClick"))
            link.set(
                qn("r:id"), part.relate_to(href, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
            )
        etree.SubElement(run, qn("a:t")).text = line


class _Deck:
    def __init__(self, title: str):
        from pptx import Presentation
        from pptx.util import Inches

        self.deck = Presentation()
        self.deck.slide_width, self.deck.slide_height = Inches(13.333), Inches(7.5)
        properties = self.deck.core_properties
        properties.author = properties.last_modified_by = "Conker"
        properties.title = title[:255]
        self.title, self.next_title = title, title
        self.current = self.body = None
        self.used, self.first = 0, True

    def slide(self, title: str):
        from pptx.util import Inches, Pt

        if len(self.deck.slides) >= MAX_SLIDES:
            raise ExportLimit("slides")
        self.current = self.deck.slides.add_slide(self.deck.slide_layouts[6])  # Blank.
        box = self.current.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(1))
        box.text_frame.word_wrap = True
        run = box.text_frame.paragraphs[0].add_run()
        run.text, run.font.size, run.font.bold = title, Pt(30), True
        self.title, self.next_title, self.body, self.used = title, None, None, 0
        return self.current

    def end(self, same_title=False):
        """The next text goes on a new slide, under the same title or marked continued."""
        base = self.title.removesuffix(CONTINUED)
        self.current, self.next_title = None, base if same_title else base + CONTINUED

    def room(self, lines: int):
        if self.current is not None and self.used and self.used + lines > LINES:
            self.end()
        if self.current is None:
            self.slide(self.next_title or self.title)
        if self.body is None:
            from pptx.enum.text import MSO_AUTO_SIZE
            from pptx.util import Inches

            box = self.current.shapes.add_textbox(
                Inches(0.6), Inches(1.45), Inches(12.1), Inches(5.6)
            )
            self.body, self.first = box.text_frame, True
            self.body.word_wrap = True
            self.body.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        self.used += lines

    def paragraph(self, spans: list[Span], block: Block, size=18, mono=False, bold=False):
        from pptx.util import Emu, Pt

        paragraph = self.body.paragraphs[0] if self.first else self.body.add_paragraph()
        self.first = False
        depth = min(block.indent + block.quote, 6)
        if depth:
            properties = paragraph._p.get_or_add_pPr()
            properties.set("marL", str(Emu(Pt(22 * depth))))
            if block.marker:
                properties.set("indent", str(-Emu(Pt(22))))
        if block.marker:
            run = paragraph.add_run()
            run.text, run.font.size = block.marker + "\t", Pt(size)
        for span in spans:
            for number, line in enumerate(span.text.split("\n")):
                if number:
                    paragraph.add_line_break()
                if not line:
                    continue
                run = paragraph.add_run()
                run.text, run.font.size = line, Pt(size)
                run.font.bold = bold or span.bold or None
                run.font.italic = span.italic or bool(block.quote) or None
                if span.code or mono:
                    run.font.name = MONO
                if span.strike:
                    run.font._rPr.set("strike", "sngStrike")
                if span.href:
                    run.hyperlink.address = span.href

    def table(self, block: Block):
        from pptx.util import Inches

        columns, size = len(block.header), 1200 if len(block.header) <= 6 else 900
        base = self.title.removesuffix(CONTINUED)
        body = block.rows or [[]]
        for start in range(0, len(body), TABLE_ROWS):
            rows = [row for row in body[start : start + TABLE_ROWS] if row]
            if self.current is None and not start:
                self.slide(self.next_title or base)
            elif start or self.body is not None or self.used:
                self.slide(base + CONTINUED)
            # Otherwise the slide holds only a heading so far, and the table goes under it.
            shape = self.current.shapes.add_table(
                len(rows) + 1,
                columns,
                Inches(0.6),
                Inches(1.45),
                Inches(12.1),
                Inches(0.4) * (len(rows) + 1),
            )
            # Cells are written as XML directly: python-pptx's cell lookup rescans the table.
            for cells, row in zip([block.header, *rows], shape.table._tbl.tr_lst, strict=True):
                for column, cell in enumerate(row.tc_lst):
                    spans = cells[column] if column < len(cells) else []
                    links = {span.href for span in spans if span.href}
                    link = links.pop() if len(links) == 1 else None
                    _cell(cell, self.current.part, plain(spans), size, link)
        self.end()


def pptx(blocks: list[Block], title: str) -> bytes:
    deck = _Deck(title)
    for block in blocks:
        if block.kind == "heading" and block.level <= 2:
            deck.slide(plain(block.spans).strip() or title)
        elif block.kind == "heading":
            deck.room(1)
            deck.paragraph(block.spans, block, size=20, bold=True)
        elif block.kind == "paragraph":
            deck.room(max(1, -(-len(plain(block.spans)) // CHARACTERS)))
            deck.paragraph(block.spans, block)
        elif block.kind == "code":
            lines = block.text.expandtabs(4).split("\n")
            for start in range(0, len(lines), CODE_LINES):
                part = lines[start : start + CODE_LINES]
                deck.room(max(1, -(-len(part) // 2)))
                deck.paragraph([Span("\n".join(part))], block, size=12, mono=True)
        elif block.kind == "rule":
            deck.end(same_title=True)
        elif block.kind == "table":
            deck.table(block)
    if not len(deck.deck.slides):
        deck.slide(title)
    output = io.BytesIO()
    deck.deck.save(output)
    return output.getvalue()
