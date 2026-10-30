"""Word documents and spreadsheets drawn from markdown blocks, with no formulas or fetches."""

from __future__ import annotations

import io
import re

from pi.content.markdown_blocks import Block, ExportLimit, Span, plain

MONO = "Courier New"
LINK = "0563C1"
STEP_CM = 0.75  # One list or quote level.
DEPTH = 8  # Deeper levels draw at this indent.


def _runs(paragraph, part, spans: list[Span], bold=False):
    """Runs written as XML directly: python-docx's run API costs ~10x more per cell."""
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml.ns import qn
    from lxml import etree

    links: dict[str, str] = {}
    for span in spans:
        parent = paragraph
        if span.href:
            if span.href not in links:
                links[span.href] = part.relate_to(
                    span.href, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
                )
            parent = etree.SubElement(paragraph, qn("w:hyperlink"))
            parent.set(qn("r:id"), links[span.href])
        # Run properties in the order the schema requires.
        looks = [
            ("w:rFonts", {"w:ascii": MONO, "w:hAnsi": MONO, "w:cs": MONO}) if span.code else None,
            ("w:b", {}) if bold or span.bold else None,
            ("w:i", {}) if span.italic else None,
            ("w:strike", {}) if span.strike else None,
            ("w:color", {"w:val": LINK}) if span.href else None,
            ("w:sz", {"w:val": "20"}) if span.code else None,
            ("w:u", {"w:val": "single"}) if span.href else None,
        ]
        run = etree.SubElement(parent, qn("w:r"))
        if any(looks):
            properties = etree.SubElement(run, qn("w:rPr"))
            for tag, values in filter(None, looks):
                element = etree.SubElement(properties, qn(tag))
                for key, value in values.items():
                    element.set(qn(key), value)
        for number, line in enumerate(span.text.split("\n")):
            if number:
                etree.SubElement(run, qn("w:br"))
            for position, piece in enumerate(line.split("\t")):
                if position:
                    etree.SubElement(run, qn("w:tab"))
                if piece:
                    text = etree.SubElement(run, qn("w:t"))
                    text.text = piece
                    text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def _layout(paragraph, block: Block):
    from docx.shared import Cm

    depth = min(block.indent + block.quote, DEPTH)
    if not depth:
        return
    form = paragraph.paragraph_format
    form.left_indent = Cm(STEP_CM * depth)
    if block.marker:
        form.first_line_indent = Cm(-STEP_CM)
        form.tab_stops.add_tab_stop(Cm(STEP_CM * depth))


def _shade(paragraph, fill="F2F2F2"):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shading)


def _rule(paragraph):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    borders, bottom = OxmlElement("w:pBdr"), OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "A6A6A6")):
        bottom.set(qn(f"w:{key}"), value)
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def _table(document, block: Block):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from lxml import etree

    table = document.add_table(rows=1 + len(block.rows), cols=len(block.header))
    table.style = "Table Grid"
    # Rows are walked once in order: python-docx's own cell lookup rescans the whole table.
    for number, row in enumerate(table._tbl.tr_lst):
        values = block.header if number == 0 else block.rows[number - 1]
        if number == 0:
            row.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for column, cell in enumerate(row.tc_lst):
            paragraph = cell.p_lst[0]
            align = block.align[column] if column < len(block.align) else None
            if align in ("left", "center", "right"):
                justify = etree.SubElement(paragraph.get_or_add_pPr(), qn("w:jc"))
                justify.set(qn("w:val"), align)
            spans = values[column] if column < len(values) else []
            _runs(paragraph, document.part, spans, bold=number == 0)


def docx(blocks: list[Block], title: str) -> bytes:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    document.core_properties.author = "Conker"
    document.core_properties.last_modified_by = "Conker"
    document.core_properties.title = title[:255]
    document.styles["Normal"].font.size = Pt(11)
    for block in blocks:
        if block.kind == "heading":
            paragraph = document.add_heading("", level=block.level)
            _runs(paragraph._p, document.part, block.spans)
        elif block.kind == "paragraph":
            paragraph = document.add_paragraph(style="Quote" if block.quote else None)
            _layout(paragraph, block)
            if block.marker:
                paragraph.add_run(block.marker + "\t")
            _runs(paragraph._p, document.part, block.spans)
        elif block.kind == "code":
            paragraph = document.add_paragraph()
            _layout(paragraph, block)
            _shade(paragraph)
            _runs(paragraph._p, document.part, [Span(block.text.expandtabs(4), code=True)])
        elif block.kind == "rule":
            _rule(document.add_paragraph())
        elif block.kind == "table":
            _table(document, block)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _sheet_name(heading: str, number: int, taken: set[str]) -> str:
    base = " ".join(re.sub(r"[\[\]:*?/\\']+", " ", heading).split())[:28].strip()
    base = base or f"Table {number}"
    name, copy = base, 2
    while name.lower() in taken:
        name, copy = f"{base[:24]} ({copy})", copy + 1
    taken.add(name.lower())
    return name


def xlsx(blocks: list[Block], title: str, sheet: str | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    tables, heading = [], ""
    for block in blocks:
        if block.kind == "heading":
            heading = plain(block.spans)
        elif block.kind == "table":
            tables.append((heading, block))
            heading = ""
    if not tables:
        raise ExportLimit("no_table")
    book = Workbook()
    book.remove(book.active)
    book.properties.creator = "Conker"
    book.properties.title = title[:255]
    taken: set[str] = set()
    for number, (heading, table) in enumerate(tables, 1):
        name = sheet if sheet and len(tables) == 1 else _sheet_name(heading, number, taken)
        page = book.create_sheet(name)
        for row_index, row in enumerate([table.header, *table.rows], 1):
            for column, spans in enumerate(row, 1):
                cell = page.cell(row_index, column)
                cell.value = plain(spans)
                cell.data_type = "s"  # Formula-looking owner/model content stays literal.
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                links = {span.href for span in spans if span.href}
                if len(links) == 1:
                    cell.hyperlink = links.pop()
                if row_index == 1:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="243B32")
        page.freeze_panes = "A2"
        page.auto_filter.ref = page.dimensions
        for column in range(1, len(table.header) + 1):
            page.column_dimensions[get_column_letter(column)].width = 24
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()
